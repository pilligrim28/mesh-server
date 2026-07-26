"""Generate Mesh Server presentation."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# Colors
BG_DARK = RGBColor(0x1A, 0x1A, 0x2E)
BG_ACCENT = RGBColor(0x16, 0x21, 0x3E)
ORANGE = RGBColor(0xFF, 0x6B, 0x35)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x2E)
GREEN = RGBColor(0x00, 0xC8, 0x53)
BLUE = RGBColor(0x00, 0x96, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape_bg(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_text_box(slide, left, top, width, height, text, font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    return tf

def add_paragraph(tf, text, font_size=18, color=WHITE, bold=False, space_before=Pt(6)):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.space_before = space_before
    return p

def add_bullet(tf, text, font_size=16, color=WHITE, level=0):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.level = level
    p.space_before = Pt(4)
    return p

# ===== SLIDE 1: Title =====
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide)

# Orange accent bar
add_shape_bg(slide, Inches(0), Inches(3.2), Inches(13.333), Inches(0.08), ORANGE)

add_text_box(slide, Inches(1), Inches(1.5), Inches(11), Inches(1.5),
             "MESH SERVER", font_size=54, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(1), Inches(3.5), Inches(11), Inches(1),
             "Сервер мониторинга устройств Meshtastic", font_size=28, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(1), Inches(5.5), Inches(11), Inches(0.5),
             "LoRa mesh-сети  |  ESP32  |  MQTT  |  Web UI", font_size=18, color=ORANGE, alignment=PP_ALIGN.CENTER)

# ===== SLIDE 2: Problem =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_shape_bg(slide, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ORANGE)

add_text_box(slide, Inches(1), Inches(0.8), Inches(11), Inches(0.8),
             "Проблема", font_size=36, color=ORANGE, bold=True)

tf = add_text_box(slide, Inches(1), Inches(2), Inches(11), Inches(5),
                  "", font_size=20, color=WHITE)
items = [
    "Нет единого решения для мониторинга mesh-устройств LoRa",
    "Сложная интеграция ESP32 с серверной логикой",
    "Отсутствие визуализации данных с носимых датчиков",
    "Нет возможности обнаруживать устройства автоматически",
    "Сложно организовать связь в полевых условиях без интернета"
]
for item in items:
    add_bullet(tf, "  " + item, font_size=20, color=WHITE)

# ===== SLIDE 3: Solution =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_shape_bg(slide, Inches(0), Inches(0), Inches(0.15), Inches(7.5), GREEN)

add_text_box(slide, Inches(1), Inches(0.8), Inches(11), Inches(0.8),
             "Решение — Mesh Server", font_size=36, color=GREEN, bold=True)

tf = add_text_box(slide, Inches(1), Inches(2), Inches(11), Inches(5),
                  "", font_size=20, color=WHITE)
items = [
    "Единая платформа для мониторинга и управления сетью",
    "Встроенный веб-интерфейс с картой и метриками",
    "Автоматическое обнаружение ESP32 через Bluetooth и WiFi",
    "MQTT-интеграция с глобальной сетью Meshtastic",
    "Работает без интернета — только LoRa mesh"
]
for item in items:
    add_bullet(tf, "  " + item, font_size=20, color=WHITE)

# ===== SLIDE 4: Features =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(1), Inches(0.5), Inches(11), Inches(0.8),
             "Возможности", font_size=36, color=ORANGE, bold=True)

features = [
    ("Карта", "GPS-координаты\nвсех устройств", ORANGE),
    ("Здоровье", "Пульс, CO2,\nТемпература", GREEN),
    ("Алерты", "Выход за зону,\nНизкий заряд", RGBColor(0xFF, 0x44, 0x44)),
    ("Сообщения", "LoRa mesh\nи MQTT", BLUE),
]

for i, (title, desc, color) in enumerate(features):
    x = Inches(1 + i * 3)
    shape = add_shape_bg(slide, x, Inches(2), Inches(2.5), Inches(2.5), BG_ACCENT)
    shape.line.color.rgb = color
    shape.line.width = Pt(2)

    add_text_box(slide, x + Inches(0.2), Inches(2.2), Inches(2.1), Inches(0.6),
                 title, font_size=22, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), Inches(3), Inches(2.1), Inches(1.2),
                 desc, font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

features2 = [
    ("Discovery", "Bluetooth\nи WiFi", RGBColor(0xAA, 0x66, 0xFF)),
    ("ML", "Аномалии и\nПредсказания", RGBColor(0xFF, 0xAA, 0x00)),
    ("WebSocket", "Realtime\nобновления", BLUE),
    ("Симулятор", "Тестовые\nданные", GREEN),
]

for i, (title, desc, color) in enumerate(features2):
    x = Inches(1 + i * 3)
    shape = add_shape_bg(slide, x, Inches(4.8), Inches(2.5), Inches(2.2), BG_ACCENT)
    shape.line.color.rgb = color
    shape.line.width = Pt(2)

    add_text_box(slide, x + Inches(0.2), Inches(5), Inches(2.1), Inches(0.6),
                 title, font_size=22, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), Inches(5.6), Inches(2.1), Inches(1),
                 desc, font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# ===== SLIDE 5: Tech specs =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(1), Inches(0.5), Inches(11), Inches(0.8),
             "Технические характеристики", font_size=36, color=ORANGE, bold=True)

specs = [
    ("Язык", "Go 1.25"),
    ("База данных", "SQLite (pure-Go)"),
    ("Протоколы", "HTTP REST, WebSocket, MQTT, BLE, USB"),
    ("Платформы", "Windows, macOS, Linux"),
    ("Интерфейсы", "Веб, Десктоп (Wails), REST API"),
    ("Размер", "~15-20 МБ"),
    ("Лицензия", "MIT (бесплатно)"),
]

for i, (key, val) in enumerate(specs):
    y = Inches(1.8 + i * 0.7)
    add_shape_bg(slide, Inches(1), y, Inches(11), Inches(0.6), BG_ACCENT if i % 2 == 0 else BG_DARK)
    add_text_box(slide, Inches(1.5), y + Inches(0.1), Inches(3), Inches(0.5),
                 key, font_size=18, color=ORANGE, bold=True)
    add_text_box(slide, Inches(5), y + Inches(0.1), Inches(6.5), Inches(0.5),
                 val, font_size=18, color=WHITE)

# ===== SLIDE 6: Advantages =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(1), Inches(0.5), Inches(11), Inches(0.8),
             "Преимущества", font_size=36, color=GREEN, bold=True)

advantages = [
    ("Все-в-одном", "Единый сервер для мониторинга, связи и анализа"),
    ("Meshtastic нативно", "Полная интеграция с публичной MQTT сетью"),
    ("LoRa без интернета", "Работает вдали от инфраструктуры"),
    ("ML на борту", "Встроенные алгоритмы аномалий"),
    ("Мульти-транспорт", "USB, Bluetooth, WiFi, MQTT"),
    ("Открытый код", "MIT лицензия — бесплатное использование"),
]

for i, (title, desc) in enumerate(advantages):
    col = i % 3
    row = i // 3
    x = Inches(1 + col * 4)
    y = Inches(2 + row * 2.5)

    add_shape_bg(slide, x, y, Inches(3.5), Inches(2), BG_ACCENT)
    add_text_box(slide, x + Inches(0.3), y + Inches(0.3), Inches(2.9), Inches(0.5),
                 title, font_size=20, color=GREEN, bold=True)
    add_text_box(slide, x + Inches(0.3), y + Inches(1), Inches(2.9), Inches(0.8),
                 desc, font_size=15, color=LIGHT_GRAY)

# ===== SLIDE 7: Architecture =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(1), Inches(0.5), Inches(11), Inches(0.8),
             "Архитектура", font_size=36, color=ORANGE, bold=True)

# Components
components = [
    ("ESP32", "LoRa + BLE", Inches(1), Inches(3)),
    ("Discovery", "Bluetooth\nWiFi Scan", Inches(4.5), Inches(2)),
    ("Mesh Server", "Go Backend", Inches(8), Inches(3)),
    ("SQLite", "Database", Inches(11.5), Inches(3)),
]

for name, desc, x, y in components:
    shape = add_shape_bg(slide, x, y, Inches(2), Inches(1.8), BG_ACCENT)
    shape.line.color.rgb = ORANGE
    shape.line.width = Pt(2)
    add_text_box(slide, x + Inches(0.1), y + Inches(0.2), Inches(1.8), Inches(0.5),
                 name, font_size=18, color=ORANGE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.1), y + Inches(0.8), Inches(1.8), Inches(0.8),
                 desc, font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# Bottom components
bottom = [
    ("Web UI", "HTML/JS", Inches(2), Inches(5.5)),
    ("Desktop", "Wails", Inches(5.5), Inches(5.5)),
    ("MQTT", "Meshtastic", Inches(9), Inches(5.5)),
]

for name, desc, x, y in bottom:
    shape = add_shape_bg(slide, x, y, Inches(2), Inches(1.5), BG_ACCENT)
    shape.line.color.rgb = BLUE
    shape.line.width = Pt(2)
    add_text_box(slide, x + Inches(0.1), y + Inches(0.2), Inches(1.8), Inches(0.5),
                 name, font_size=16, color=BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.1), y + Inches(0.7), Inches(1.8), Inches(0.5),
                 desc, font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# ===== SLIDE 8: Readiness =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(1), Inches(0.5), Inches(11), Inches(0.8),
             "Стадия готовности", font_size=36, color=ORANGE, bold=True)

add_text_box(slide, Inches(1), Inches(1.8), Inches(11), Inches(0.8),
             "Готов к демонстрации и тестированию", font_size=24, color=GREEN, bold=True)

items = [
    ("Веб-интерфейс", True),
    ("Десктопное приложение (Wails)", True),
    ("REST API (полный набор)", True),
    ("MQTT интеграция", True),
    ("Bluetooth/WiFi Discovery", True),
    ("ML компоненты", True),
    ("Документация", True),
    ("Демо-режим", True),
]

tf = add_text_box(slide, Inches(1), Inches(3), Inches(11), Inches(4),
                  "", font_size=18, color=WHITE)
for item, ready in items:
    status = "Готово" if ready else "В разработке"
    color = GREEN if ready else RGBColor(0xFF, 0xAA, 0x00)
    add_bullet(tf, f"  {item} — {status}", font_size=18, color=color)

# ===== SLIDE 9: End =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_shape_bg(slide, Inches(0), Inches(3.2), Inches(13.333), Inches(0.08), ORANGE)

add_text_box(slide, Inches(1), Inches(2), Inches(11), Inches(1),
             "Спасибо за внимание!", font_size=48, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(1), Inches(3.8), Inches(11), Inches(0.8),
             "Mesh Server — ваш мониторинг mesh-сетей", font_size=24, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(1), Inches(5.5), Inches(11), Inches(0.5),
             "GitHub: open source  |  MIT License", font_size=18, color=ORANGE, alignment=PP_ALIGN.CENTER)

# Save
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "MeshServer_Presentation.pptx")
prs.save(output_path)
print(f"Saved: {output_path}")
